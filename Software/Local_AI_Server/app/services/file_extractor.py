import io
import os
from typing import Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None


def extract_text_from_file_bytes(
    file_bytes: bytes,
    filename: str,
    content_type: str = "",
) -> str:
    """
    Extract clean plain text from PDF, DOCX, or TXT file bytes.
    """
    lower_name = (filename or "").lower()

    # 1. Plain Text / CSV
    if lower_name.endswith((".txt", ".csv")) or "text/" in content_type:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")

    # 2. PDF Document
    if lower_name.endswith(".pdf") or "pdf" in content_type:
        if fitz is None:
            raise RuntimeError("PyMuPDF (fitz) is not installed on server.")

        extracted_pages = []
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                page_text = page.get_text("text") or ""
                if page_text.strip():
                    extracted_pages.append(page_text.strip())

        combined = "\n\n".join(extracted_pages).strip()
        if not combined:
            return "[Cảnh báo: File PDF không có text layer hoặc là tài liệu scan dạng hình ảnh.]"
        return combined

    # 3. Microsoft Word DOCX
    if lower_name.endswith(".docx") or "officedocument" in content_type:
        if docx is None:
            raise RuntimeError("python-docx is not installed on server.")

        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n\n".join(paragraphs).strip()

    # Fallback to UTF-8 decoding
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="replace")
